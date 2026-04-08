#!/usr/bin/env python3
"""Process raw data files into clean JSON for WA breweries, wineries, and distilleries."""

import json
import openpyxl
import os

BASE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(BASE, "raw")

###############
# BREWERIES - from Open Brewery DB
###############
all_breweries = []
for page in [1, 2, 3]:
    fpath = os.path.join(RAW, f"openbrewerydb_wa_page{page}.json")
    with open(fpath) as f:
        data = json.load(f)
        for b in data:
            all_breweries.append({
                "name": b.get("name", "").strip(),
                "city": (b.get("city") or "").strip(),
                "website": b.get("website_url") or "",
                "brewery_type": b.get("brewery_type", ""),
                "address": (b.get("street") or "").strip(),
                "phone": b.get("phone") or ""
            })

seen = set()
deduped = []
for b in all_breweries:
    key = b["name"].lower()
    if key not in seen:
        seen.add(key)
        deduped.append(b)

brewery_json = {
    "source": "Open Brewery DB API",
    "source_url": "https://api.openbrewerydb.org/v1/breweries?by_state=washington&per_page=200",
    "extracted_date": "2026-04-07",
    "type": "brewery",
    "count": len(deduped),
    "entries": sorted(deduped, key=lambda x: x["name"].lower())
}

with open(os.path.join(BASE, "wa_breweries.json"), "w") as f:
    json.dump(brewery_json, f, indent=2)
print(f"Breweries: {len(deduped)} entries saved")

###############
# WINERIES - from LCB Active Winery List
###############
wb = openpyxl.load_workbook(os.path.join(RAW, "lcb_active_wineries.xlsx"))
ws = wb.active

wineries = []
seen_wineries = set()
for row in ws.iter_rows(min_row=7, values_only=True):
    vals = list(row)
    license_num = str(vals[0]).strip() if vals[0] else ""
    if license_num and license_num.isdigit() and len(license_num) >= 5:
        name = (vals[1] or "").strip()
        # Winery columns: 0=License#, 1=Licensee, 3=Address, 6=City, 8=State
        city = (vals[6] or "").strip() if len(vals) > 6 and vals[6] else ""
        if name:
            name_title = name.strip().title()
            city_title = city.strip().title()
            key = name.strip().lower()
            if key not in seen_wineries:
                seen_wineries.add(key)
                wineries.append({
                    "name": name_title,
                    "city": city_title,
                    "license_number": license_num,
                    "website": ""
                })

winery_json = {
    "source": "WA Liquor and Cannabis Board - Active Winery Licensees",
    "source_url": "https://lcb.wa.gov/sites/default/files/2025-09/Active%20Winery%20List%2009172025.xlsx",
    "extracted_date": "2026-04-07",
    "type": "winery",
    "count": len(wineries),
    "entries": sorted(wineries, key=lambda x: x["name"].lower())
}

with open(os.path.join(BASE, "wa_wineries.json"), "w") as f:
    json.dump(winery_json, f, indent=2)
print(f"Wineries: {len(wineries)} entries saved")

###############
# DISTILLERIES - from all 3 LCB files combined
###############
distilleries = []
seen_dist = set()

for fname, source_type in [
    ("lcb_craft_distillery.xlsx", "Craft Distillery"),
    ("lcb_distiller_rectifier.xlsx", "Distiller/Rectifier"),
    ("lcb_fruit_distiller.xlsx", "Fruit/Wine Distillery")
]:
    wb = openpyxl.load_workbook(os.path.join(RAW, fname))
    ws = wb.active
    for row in ws.iter_rows(min_row=6, values_only=True):
        vals = list(row)
        # Columns: 0=Privilege, 1=Tradename, 2=(empty), 3=License#, 4=(empty),
        #          5=UBI, 6=Licensee, 7=Status, 8=Location Address, 9=room#, 10=City, 11=St
        privilege = (str(vals[0]) if vals[0] else "").strip().upper()
        tradename = (str(vals[1]) if vals[1] else "").strip()
        license_num = (str(vals[3]) if len(vals) > 3 and vals[3] else "").strip()
        status = (str(vals[7]) if len(vals) > 7 and vals[7] else "").strip()
        city = (str(vals[10]) if len(vals) > 10 and vals[10] else "").strip()

        # Only include distillery privileges
        if not any(kw in privilege for kw in ["DISTILL", "CRAFT DISTILLERY", "FRUIT"]):
            continue

        if tradename:
            name_title = tradename.strip().title()
            key = tradename.strip().lower()
            if key not in seen_dist:
                seen_dist.add(key)
                distilleries.append({
                    "name": name_title,
                    "city": city.strip().title(),
                    "license_number": license_num,
                    "license_type": source_type,
                    "status": status,
                    "website": ""
                })

distillery_json = {
    "source": "WA Liquor and Cannabis Board - Craft Distillery, Distiller/Rectifier, and Fruit Distiller Lists",
    "source_url": "https://lcb.wa.gov/records/frequently-requested-lists",
    "extracted_date": "2026-04-07",
    "type": "distillery",
    "count": len(distilleries),
    "entries": sorted(distilleries, key=lambda x: x["name"].lower())
}

with open(os.path.join(BASE, "wa_distilleries.json"), "w") as f:
    json.dump(distillery_json, f, indent=2)

active_dist = [d for d in distilleries if "ACTIVE" in d["status"].upper()]
expired_dist = [d for d in distilleries if "EXPIRED" in d["status"].upper()]
print(f"Distilleries: {len(distilleries)} entries saved (Active: {len(active_dist)}, Expired: {len(expired_dist)})")
